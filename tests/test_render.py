from pathlib import Path

import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches

from tests.test_parse import SAMPLE
from tools.render_cv import (
    SofficeNotFound,
    find_dashes,
    find_soffice,
    parse_cv,
    render_docx,
)


@pytest.fixture
def doc():
    return render_docx(parse_cv(SAMPLE))


def _para(doc, text):
    return next(p for p in doc.paragraphs if p.text.startswith(text))


def test_name_is_the_first_paragraph_in_bold_17pt(doc):
    run = doc.paragraphs[0].runs[0]
    assert (run.text, run.bold, run.font.size.pt) == ("Ana Pérez", True, 17)


def test_section_heading_is_centered_bold_with_bottom_border(doc):
    p = _para(doc, "Experiencia")
    assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert p.runs[0].bold
    assert p._p.pPr.find(qn("w:pBdr")) is not None


def test_entry_line_puts_dates_on_a_right_tab_stop(doc):
    p = _para(doc, "Banco Ejemplo")
    assert p.text == "Banco Ejemplo\tMar 2024 a Dic 2025"
    assert p.paragraph_format.tab_stops[0].position == Inches(7.1)


def test_contact_links_are_real_hyperlinks(doc):
    links = doc.paragraphs[2]._p.findall(qn("w:hyperlink"))
    targets = {doc.part.rels[h.get(qn("r:id"))].target_ref for h in links}
    assert targets == {"mailto:ana@example.com", "https://linkedin.com/in/ana-example"}


def test_bullets_are_a_dash_with_hanging_indent(doc):
    p = _para(doc, "-\tConstruí")
    assert round(p.paragraph_format.left_indent.inches, 2) == 0.22
    assert round(p.paragraph_format.first_line_indent.inches, 2) == -0.22


def test_project_line_shows_title_and_stack(doc):
    assert (
        _para(doc, "Valuación DCF").text
        == "Valuación DCF de una fintech · Excel, Python"
    )


def test_find_dashes_reports_line_numbers_with_em_or_en_dash():
    assert find_dashes("ok\nmal — aca\nbien\nmal – tambien") == [2, 4]


def test_find_dashes_ignores_hyphens():
    assert find_dashes("2024-2025 y bien") == []


def test_find_soffice_uses_windows_default_install_path(tmp_path):
    exe = tmp_path / "LibreOffice" / "program" / "soffice.exe"
    exe.parent.mkdir(parents=True)
    exe.touch()
    assert find_soffice(which=lambda _: None, candidates=[exe]) == exe


def test_find_soffice_raises_plain_language_error_when_missing():
    with pytest.raises(SofficeNotFound, match="LibreOffice"):
        find_soffice(which=lambda _: None, candidates=[Path("/nope/soffice.exe")])


def test_find_dashes_ignores_dashes_inside_html_comments():
    assert find_dashes("bien\n<!--\n2025–2026\n-->\nbien") == []


from tools.render_cv import application_copy  # noqa: E402


def test_application_copy_uses_the_neutral_name_in_a_per_application_folder(tmp_path):
    pdf = tmp_path / "banco-analista.pdf"
    pdf.write_bytes(b"%PDF")
    out = application_copy(pdf, tmp_path / "para-postular", "Banco Macro_Analista", "CV_Ana_Perez")
    assert out == tmp_path / "para-postular" / "Banco Macro_Analista" / "CV_Ana_Perez.pdf"
    assert out.read_bytes() == b"%PDF"


def test_english_cv_gets_the_en_suffix(tmp_path):
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF")
    assert application_copy(pdf, tmp_path, "JP Morgan_Associate", "CV_Ana_Perez", english=True).name == "CV_Ana_Perez_EN.pdf"


def test_folder_name_drops_characters_windows_rejects(tmp_path):
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"%PDF")
    out = application_copy(pdf, tmp_path, 'Ualá: "Business" Control/Lead?', "CV_Ana_Perez")
    assert out.parent.name == "Ualá Business ControlLead"


def test_existing_copy_is_replaced(tmp_path):
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(b"v2")
    target = tmp_path / "A_B" / "CV_Ana_Perez.pdf"
    target.parent.mkdir()
    target.write_bytes(b"v1")
    assert application_copy(pdf, tmp_path, "A_B", "CV_Ana_Perez").read_bytes() == b"v2"

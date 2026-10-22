"""Render a CV written in the strict markdown format (see docs/FORMATO-CV.md) to .docx and .pdf."""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.shared import OxmlElement, qn
from docx.shared import Inches, Pt

LINK_RE = re.compile(r"^\[(?P<text>[^\]]+)\]\((?P<url>[^)]+)\)$")
SKILL_RE = re.compile(r"^\*\*(?P<label>[^*]+)\*\*\s*(?P<text>.*)$")
ITALIC_RE = re.compile(r"^_(?P<text>.+)_$")
COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)


def _strip_comments(text: str) -> str:
    """Drop HTML comments (e.g. the saved job posting) but keep their newlines so line numbers stay valid."""
    return COMMENT_RE.sub(lambda m: "\n" * m.group().count("\n"), text)


@dataclass
class Block:
    kind: str  # "paragraph" | "entry" | "project" | "skill" | "bullets"
    text: str = ""
    title: str = ""
    dates: str = ""
    role: str = ""
    location: str = ""
    stack: str = ""
    context: str = ""
    label: str = ""
    bullets: list[str] = field(default_factory=list)


@dataclass
class Section:
    title: str
    blocks: list[Block] = field(default_factory=list)


@dataclass
class CV:
    name: str = ""
    tagline: str = ""
    contact: list[tuple[str, str | None]] = field(default_factory=list)
    sections: list[Section] = field(default_factory=list)


def _parse_contact(line: str) -> list[tuple[str, str | None]]:
    segments = []
    for raw in line.split(" | "):
        raw = raw.strip()
        m = LINK_RE.match(raw)
        segments.append((m["text"], m["url"]) if m else (raw, None))
    return segments


def _parse_heading(heading: str) -> Block:
    if " | " in heading:
        title, dates = heading.split(" | ", 1)
        return Block(kind="entry", title=title.strip(), dates=dates.strip())
    if " · " in heading:
        title, stack = heading.split(" · ", 1)
        return Block(kind="project", title=title.strip(), stack=stack.strip())
    return Block(kind="entry", title=heading.strip())


def _parse_body_line(section: Section, line: str) -> None:
    last = section.blocks[-1] if section.blocks else None
    if line.startswith("### "):
        section.blocks.append(_parse_heading(line[4:]))
    elif line.startswith("- "):
        if last is None or last.kind not in ("entry", "project", "bullets"):
            last = Block(kind="bullets")
            section.blocks.append(last)
        last.bullets.append(line[2:].strip())
    elif (
        (m := ITALIC_RE.match(line))
        and last
        and last.kind in ("entry", "project")
        and not last.bullets
    ):
        if last.kind == "entry":
            role, _, location = m["text"].partition(" | ")
            last.role, last.location = role.strip(), location.strip()
        else:
            last.context = m["text"].strip()
    elif m := SKILL_RE.match(line):
        section.blocks.append(
            Block(kind="skill", label=m["label"].strip(), text=m["text"].strip())
        )
    else:
        section.blocks.append(Block(kind="paragraph", text=line))


def parse_cv(text: str) -> CV:
    cv = CV()
    header_lines: list[str] = []
    section: Section | None = None
    for line in _strip_comments(text).splitlines():
        stripped = line.strip()
        if line.startswith("# "):
            cv.name = line[2:].strip()
        elif line.startswith("## "):
            section = Section(title=line[3:].strip())
            cv.sections.append(section)
        elif not stripped or stripped == "---":
            continue
        elif section is None:
            header_lines.append(stripped)
        else:
            _parse_body_line(section, stripped)
    if header_lines:
        cv.tagline = header_lines[0]
    if len(header_lines) > 1:
        cv.contact = _parse_contact(header_lines[1])
    return cv


# ---------- rendering ----------


RIGHT_TAB = Inches(7.1)  # 8.5in page - 0.7in side margins
LINK_COLOR = "1155CC"
HANGING = Inches(0.22)


def _set_bottom_border(p) -> None:
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "2"), ("w:color", "000000")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    p_pr.append(borders)


def _add_hyperlink(p, url: str, text: str) -> None:
    r_id = p.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi"):
        fonts.set(qn(attr), "Arial")
    r_pr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    p._p.append(link)


def render_docx(cv: CV):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.55)
    sec.left_margin = sec.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    def para(before=0, after=0, align=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        if align is not None:
            p.alignment = align
        return p

    def tabbed(p, left: str, right: str, bold_left=False) -> None:
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT_TAB, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(left).bold = bold_left
        if right:
            p.add_run("\t" + right)

    def bullet(text: str) -> None:
        p = para(after=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.left_indent = HANGING
        p.paragraph_format.first_line_indent = -HANGING
        p.paragraph_format.tab_stops.add_tab_stop(HANGING, WD_TAB_ALIGNMENT.LEFT)
        p.add_run("-\t" + text)

    name = para().add_run(cv.name)
    name.bold = True
    name.font.size = Pt(17)
    para(after=2).add_run(cv.tagline)
    contact = para(after=4)
    for i, (text, url) in enumerate(cv.contact):
        if i:
            contact.add_run(" | ")
        if url:
            _add_hyperlink(contact, url, text)
        else:
            contact.add_run(text)

    for section in cv.sections:
        heading = para(before=7, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = heading.add_run(section.title)
        run.bold = True
        run.font.size = Pt(10.5)
        _set_bottom_border(heading)
        for block in section.blocks:
            if block.kind == "paragraph":
                para(after=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY).add_run(block.text)
            elif block.kind == "skill":
                p = para(after=2)
                p.add_run(block.label + " ").bold = True
                p.add_run(block.text)
            elif block.kind == "entry":
                tabbed(para(before=5), block.title, block.dates, bold_left=True)
                if block.role or block.location:
                    tabbed(para(after=3), block.role, block.location)
            elif block.kind == "project":
                p = para(before=6)
                p.add_run(block.title).bold = True
                if block.stack:
                    p.add_run(" · " + block.stack)
                if block.context:
                    ctx = para(after=2).add_run(block.context)
                    ctx.italic = True
                    ctx.font.size = Pt(9)
            for b in block.bullets:
                bullet(b)
    return doc


def find_dashes(text: str) -> list[int]:
    """Line numbers (1-based) containing an em dash or en dash, which Spanish CVs must not use."""
    return [n for n, line in enumerate(_strip_comments(text).splitlines(), 1) if "\u2014" in line or "\u2013" in line]


class SofficeNotFound(RuntimeError):
    pass


WINDOWS_SOFFICE = [
    Path(r"C:\Program Files\LibreOffice\program\soffice.exe"),
    Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"),
]


def find_soffice(which=shutil.which, candidates=WINDOWS_SOFFICE) -> Path:
    for name in ("soffice", "libreoffice"):
        if found := which(name):
            return Path(found)
    for path in candidates:
        if path.exists():
            return path
    raise SofficeNotFound(
        "No encontré LibreOffice, que se usa para generar el PDF. "
        "Instalalo desde https://www.libreoffice.org/download/ y volvé a intentar. "
        "El .docx ya quedó generado."
    )


class PdfConversionError(RuntimeError):
    pass


def _word_convert(docx_path: Path) -> None:
    """Convert with Microsoft Word (Windows only, via docx2pdf). Word is the reference renderer for .docx."""
    if sys.platform != "win32":
        raise RuntimeError("Word is only used on Windows")
    from docx2pdf import convert

    convert(str(docx_path), str(docx_path.with_suffix(".pdf")))


def _soffice_convert(docx_path: Path) -> None:
    subprocess.run(
        [str(find_soffice()), "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
        check=True,
        capture_output=True,
    )


def convert_to_pdf(docx_path: Path, backends=(_word_convert, _soffice_convert)) -> Path:
    """Try each backend in order; the first one that actually produces the PDF wins."""
    pdf_path = docx_path.with_suffix(".pdf")
    for backend in backends:
        pdf_path.unlink(missing_ok=True)
        try:
            backend(docx_path)
        except Exception:
            continue
        if pdf_path.exists():
            return pdf_path
    raise PdfConversionError(
        "No pude generar el PDF: no funcionó ni Word ni LibreOffice. "
        "Revisá que Word abra bien en esta computadora. El .docx ya quedó generado."
    )


WINDOWS_FORBIDDEN = re.compile(r'[<>:"/\\|?*]')


def application_copy(pdf: Path, base: Path, empresa_puesto: str, name: str, english: bool = False) -> Path:
    """Copy the PDF to <base>/<Empresa_Puesto>/<name>[_EN].pdf: the file the user sends never names the company."""
    folder = base / " ".join(WINDOWS_FORBIDDEN.sub("", empresa_puesto).split())
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / f"{name}{'_EN' if english else ''}.pdf"
    shutil.copyfile(pdf, target)
    return target


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Genera .docx y .pdf a partir de un CV en markdown.")
    parser.add_argument("cv_md", type=Path)
    parser.add_argument("--no-pdf", action="store_true")
    parser.add_argument("--postulacion", help="Empresa_Puesto: also copy the PDF with the neutral name (pdf_nombre, pdf_carpeta in profile/busqueda.toml)")
    parser.add_argument("--en", action="store_true", help="the CV is in English (adds _EN to the neutral name)")
    args = parser.parse_args(argv)

    text = args.cv_md.read_text(encoding="utf-8")
    if lines := find_dashes(text):
        print(f"El CV tiene rayas largas en las líneas {lines}. Reemplazalas antes de generar.", file=sys.stderr)
        return 1
    docx_path = args.cv_md.with_suffix(".docx")
    render_docx(parse_cv(text)).save(str(docx_path))
    print(f"Listo: {docx_path}")
    if not args.no_pdf:
        try:
            pdf_path = convert_to_pdf(docx_path)
            print(f"Listo: {pdf_path}")
            if args.postulacion:
                import tomllib

                cfg = tomllib.loads(Path("profile/busqueda.toml").read_text(encoding="utf-8"))
                if cfg.get("pdf_nombre") and cfg.get("pdf_carpeta"):
                    copy = application_copy(pdf_path, Path(cfg["pdf_carpeta"]).expanduser(), args.postulacion, cfg["pdf_nombre"], english=args.en)
                    print(f"Para postular: {copy}")
                else:
                    print("Falta pdf_nombre o pdf_carpeta en profile/busqueda.toml: no copié la versión con nombre neutro.", file=sys.stderr)
        except PdfConversionError as exc:
            print(exc, file=sys.stderr)
            return 2
    return 0


if __name__ == "__main__":
    sys.exit(main())

"""Google Sheets tracker: one .xlsx file inside their Google Drive folder, which Google Sheets opens on any device.

  uv run python -m tools.planilla crear <ruta.xlsx>              # create the tracker (columns, Estado dropdown)
  uv run python -m tools.planilla leer <ruta.xlsx>               # print every row as JSON
  uv run python -m tools.planilla guardar <ruta.xlsx> <filas.json>

`guardar` takes a JSON list of rows ({"Columna": valor}). A row whose `Link` (or `Puesto` + `Empresa` when there is no
link) is already in the file is updated with the given columns only; any other row is appended. The user's own
changes (for example a new `Estado`) are kept, because only the columns in the JSON are written.

Google Drive for desktop syncs the folder both ways, so the user sees and edits the same file on their phone.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HOJA = "Postulaciones"
COLUMNAS = [
    "Puesto", "Empresa", "Estado", "Puntaje", "Motivo", "Sector", "Modalidad", "Exterior", "Link", "Plataforma",
    "CV", "Próximo paso", "Fecha", "Orden", "Descarte automático", "Fecha postulación", "Confirmación recibida",
    "Requiere acción", "Email", "Última novedad",
]
ESTADOS = [
    "Nuevo", "Para decidir", "Descartado", "Aprobado", "CV listo", "Postulado", "En proceso", "Entrevista",
    "Oferta", "Rechazado", "Sin respuesta", "Retirado",
]
ANCHOS = {"Puesto": 40, "Empresa": 22, "Estado": 15, "Motivo": 50, "Link": 40, "Próximo paso": 35, "Última novedad": 40}


def crear(ruta: Path) -> None:
    if ruta.exists():
        raise SystemExit(f"Ya existe {ruta}; no la piso.")
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA
    ws.append(COLUMNAS)
    for i, nombre in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F3A5F")
        ws.column_dimensions[celda.column_letter].width = ANCHOS.get(nombre, 16)
    ws.freeze_panes = "A2"
    col = ws.cell(row=1, column=COLUMNAS.index("Estado") + 1).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(ESTADOS) + '"', allow_blank=True)
    dv.add(f"{col}2:{col}5000")
    ws.add_data_validation(dv)
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(COLUMNAS)).column_letter}1"
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


def _encabezados(ws) -> list[str]:
    return [c.value for c in ws[1]]


def leer(ruta: Path) -> list[dict]:
    ws = load_workbook(ruta)[HOJA]
    cols = _encabezados(ws)
    filas = []
    for valores in ws.iter_rows(min_row=2, values_only=True):
        if any(v not in (None, "") for v in valores):
            filas.append({c: (v.isoformat() if hasattr(v, "isoformat") else v) for c, v in zip(cols, valores) if c})
    return filas


def _clave(fila: dict) -> str:
    if fila.get("Link"):
        return "link:" + str(fila["Link"]).strip().rstrip("/")
    return "pe:" + str(fila.get("Puesto", "")).strip().lower() + "|" + str(fila.get("Empresa", "")).strip().lower()


def guardar(ruta: Path, nuevas: list[dict]) -> tuple[int, int]:
    wb = load_workbook(ruta)
    ws = wb[HOJA]
    cols = _encabezados(ws)
    desconocidas = {k for f in nuevas for k in f} - set(cols)
    if desconocidas:
        raise SystemExit(f"Columnas que no existen en la planilla: {', '.join(sorted(desconocidas))}")
    indice = {}
    for n, valores in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        indice.setdefault(_clave(dict(zip(cols, valores))), n)
    actualizadas = agregadas = 0
    for fila in nuevas:
        n = indice.get(_clave(fila))
        if n is None:
            n = ws.max_row + 1
            indice[_clave(fila)] = n
            agregadas += 1
        else:
            actualizadas += 1
        for k, v in fila.items():
            ws.cell(row=n, column=cols.index(k) + 1, value=v)
    wb.save(ruta)
    return actualizadas, agregadas


def main(argv: list[str]) -> None:
    if len(argv) < 2 or argv[0] not in {"crear", "leer", "guardar"}:
        raise SystemExit(__doc__)
    ruta = Path(argv[1])
    if argv[0] == "crear":
        crear(ruta)
        print(f"Planilla creada: {ruta}")
    elif argv[0] == "leer":
        sys.stdout.reconfigure(encoding="utf-8")
        print(json.dumps(leer(ruta), ensure_ascii=False, indent=1))
    else:
        filas = json.loads(Path(argv[2]).read_text(encoding="utf-8"))
        act, agr = guardar(ruta, filas)
        print(f"Actualizadas: {act}. Agregadas: {agr}.")


if __name__ == "__main__":
    main(sys.argv[1:])
